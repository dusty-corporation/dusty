"""Render Dusty transcript formats with controlled line endings and metadata."""

import io
import json
import re
import textwrap
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from dusty.core import Segment, Transcript


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _paragraphs(segments: tuple[Segment, ...]) -> list[tuple[float, str]]:
    paragraphs: list[tuple[float, str]] = []
    current: list[str] = []
    first_start = 0.0
    previous_end = 0.0
    for segment in segments:
        text = _clean(segment.text)
        if not text:
            continue
        if current and segment.start - previous_end >= 2.0:
            paragraphs.append((first_start, " ".join(current)))
            current = []
        if not current:
            first_start = segment.start
        current.append(text)
        previous_end = segment.end
    if current:
        paragraphs.append((first_start, " ".join(current)))
    return paragraphs


def render_txt(transcript: Transcript) -> bytes:
    paragraphs = _paragraphs(transcript.segments)
    if not paragraphs:
        return b""
    return ("\r\n\r\n".join(text for _, text in paragraphs) + "\r\n").encode("utf-8")


def _timestamp(seconds: float, *, separator: str) -> str:
    milliseconds = max(0, round(seconds * 1000))
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, millis = divmod(remainder, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02}{separator}{millis:03}"


def _segment_cues(segment: Segment) -> list[tuple[float, float, str]]:
    text = _clean(segment.text)
    if not text:
        return []
    lines = textwrap.wrap(text, width=42, break_long_words=True, break_on_hyphens=False)
    chunks = ["\n".join(lines[index : index + 2]) for index in range(0, len(lines), 2)]
    total_length = sum(len(chunk.replace("\n", "")) for chunk in chunks)
    duration = max(0.0, segment.end - segment.start)
    cues: list[tuple[float, float, str]] = []
    consumed = 0
    for chunk in chunks:
        start = segment.start + duration * consumed / total_length
        consumed += len(chunk.replace("\n", ""))
        end = segment.start + duration * consumed / total_length
        cues.append((start, end, chunk))
    return cues


def _captions(transcript: Transcript, *, vtt: bool) -> bytes:
    separator = "." if vtt else ","
    lines = ["WEBVTT", ""] if vtt else []
    cue_index = 0
    for segment in transcript.segments:
        for start, end, text in _segment_cues(segment):
            cue_index += 1
            if not vtt:
                lines.append(str(cue_index))
            lines.extend(
                [
                    f"{_timestamp(start, separator=separator)} --> "
                    f"{_timestamp(end, separator=separator)}",
                    text,
                    "",
                ]
            )
    return "\r\n".join(lines).encode("utf-8")


def render_srt(transcript: Transcript) -> bytes:
    return _captions(transcript, vtt=False)


def render_vtt(transcript: Transcript) -> bytes:
    return _captions(transcript, vtt=True)


def render_json(
    transcript: Transcript, source: Path, *, created_utc: datetime | None = None
) -> bytes:
    created = created_utc or datetime.now(UTC)
    value = {
        "schema": "dusty.transcript",
        "schema_version": 1,
        "app_version": "0.0.1",
        "model_id": transcript.model_id,
        "backend": transcript.backend,
        "task": transcript.task,
        "language": transcript.language,
        "duration_sec": transcript.duration_sec,
        "source_file_name": source.name,
        "created_utc": created.astimezone(UTC).isoformat().replace("+00:00", "Z"),
        "segments": [
            {"index": part.index, "start": part.start, "end": part.end, "text": _clean(part.text)}
            for part in transcript.segments
            if _clean(part.text)
        ],
    }
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def render_docx(
    transcript: Transcript,
    source: Path,
    *,
    created_utc: datetime | None = None,
    timestamps: bool = True,
) -> bytes:
    document = Document()
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.identifier = ""
    properties.subject = ""
    properties.title = ""
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    document.add_heading(source.stem, level=0)
    created = created_utc or datetime.now(UTC)
    date = created.astimezone(UTC).strftime("%Y-%m-%d")
    document.add_paragraph(
        f"{date} | {transcript.duration_sec:.1f} s | {transcript.model_id} | {transcript.language}"
    )
    for start, text in _paragraphs(transcript.segments):
        prefix = f"[{_timestamp(start, separator='.').split('.')[0]}] " if timestamps else ""
        document.add_paragraph(prefix + text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
