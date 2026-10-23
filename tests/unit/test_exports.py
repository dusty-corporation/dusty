"""Check transcript format boundaries, metadata, and caption limits."""

import io
import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document

from dusty.core import Segment, Transcript
from dusty.export import render_docx, render_json, render_srt, render_txt, render_vtt


def sample_transcript() -> Transcript:
    return Transcript(
        segments=(
            Segment(0, 0.0, 1.0, "  Hello   world.  "),
            Segment(1, 1.2, 2.0, "Another line."),
            Segment(2, 4.0, 7.0, "This is a later paragraph."),
        ),
        language="en",
        task="transcribe",
        model_id="tiny",
        backend="cpu",
        duration_sec=7.0,
    )


def test_txt_paragraph_gap_and_empty_transcript() -> None:
    transcript = sample_transcript()
    assert render_txt(transcript) == (
        b"Hello world. Another line.\r\n\r\nThis is a later paragraph.\r\n"
    )
    assert render_txt(Transcript((), "en", "transcribe", "tiny", "cpu", 0)) == b""


def test_captions_split_long_text_into_short_lines() -> None:
    words = " ".join(["transcription"] * 20)
    transcript = Transcript((Segment(0, 0, 10, words),), "en", "transcribe", "tiny", "cpu", 10)
    for render in (render_srt, render_vtt):
        content = render(transcript).decode("utf-8")
        assert "00:00:00" in content
        assert "00:00:10" in content
        text_lines = [line for line in content.splitlines() if "transcription" in line]
        assert len(text_lines) > 2
        assert all(len(line) <= 42 for line in text_lines)
    assert render_vtt(transcript).startswith(b"WEBVTT\r\n\r\n")
    assert b"00:00:00,000 -->" in render_srt(transcript)


def test_json_uses_basename_only_and_no_hidden_path() -> None:
    payload = render_json(
        sample_transcript(),
        Path("C:/private/recording.wav"),
        created_utc=datetime(2026, 10, 22, tzinfo=UTC),
    )
    output = json.loads(payload)
    assert output["source_file_name"] == "recording.wav"
    assert output["created_utc"] == "2026-10-22T00:00:00Z"
    assert "C:/private" not in payload.decode("utf-8")
    assert len(output["segments"]) == 3


def test_docx_clears_author_metadata() -> None:
    payload = render_docx(sample_transcript(), Path("recording.wav"))
    document = Document(io.BytesIO(payload))
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""
    assert document.paragraphs[0].text == "recording"
    assert any("Hello world." in paragraph.text for paragraph in document.paragraphs)


def test_stdlib_docx_opens_with_escaped_unicode_and_no_timestamps() -> None:
    transcript = Transcript(
        (Segment(0, 0, 2, '<tag> & "سلام" 😀\x01'),), "fa", "transcribe", "tiny", "cpu", 2
    )
    payload = render_docx(transcript, Path("A & B.wav"), timestamps=False)
    document = Document(io.BytesIO(payload))
    assert document.paragraphs[0].text == "A & B"
    assert document.paragraphs[-1].text == '<tag> & "سلام" 😀'
    for name in (
        "author",
        "last_modified_by",
        "comments",
        "title",
        "subject",
        "keywords",
        "category",
        "identifier",
    ):
        assert getattr(document.core_properties, name) == ""


def golden_transcript() -> Transcript:
    return Transcript(
        (Segment(0, 0, 1, "  Hello  world. "), Segment(1, 3, 4, "سلام دنیا")),
        "en",
        "transcribe",
        "tiny",
        "cpu",
        4,
    )


@pytest.mark.parametrize("format_name", ["txt", "srt", "vtt"])
def test_text_format_goldens(format_name: str) -> None:
    root = Path(__file__).resolve().parents[1] / "fixtures/exports"
    render = {"txt": render_txt, "srt": render_srt, "vtt": render_vtt}[format_name]
    expected = (root / f"sample.{format_name}").read_text(encoding="utf-8")
    assert render(golden_transcript()) == expected.replace("\n", "\r\n").encode("utf-8")


def test_json_and_docx_semantic_goldens() -> None:
    root = Path(__file__).resolve().parents[1] / "fixtures/exports"
    created = datetime(2026, 10, 22, tzinfo=UTC)
    transcript = golden_transcript()
    payload = render_json(transcript, Path("sample.wav"), created_utc=created)
    assert json.loads(payload) == json.loads((root / "sample.json").read_text(encoding="utf-8"))
    expected = json.loads((root / "sample-docx.json").read_text(encoding="utf-8"))
    document = Document(
        io.BytesIO(render_docx(transcript, Path("sample.wav"), created_utc=created))
    )
    assert [paragraph.text for paragraph in document.paragraphs] == expected["paragraphs"]
    font = document.styles["Normal"].font
    assert font.name == expected["font"]
    assert font.size.pt == expected["font_size_pt"]
    for property_name in ("author", "last_modified_by", "comments"):
        assert getattr(document.core_properties, property_name) == expected[property_name]


def test_wrapped_captions_have_only_crlf_and_two_lines_per_cue() -> None:
    transcript = Transcript(
        (Segment(0, 3599.9996, 3609, " ".join(["سلام"] * 80)),),
        "fa",
        "transcribe",
        "tiny",
        "cpu",
        3610,
    )
    for render in (render_srt, render_vtt):
        payload = render(transcript)
        assert b"\n" not in payload.replace(b"\r\n", b"")
        assert b"01:00:00" in payload
        for cue in payload.decode("utf-8").strip().split("\r\n\r\n"):
            lines = cue.split("\r\n")
            timing = next((index for index, line in enumerate(lines) if " --> " in line), None)
            if timing is not None:
                assert 1 <= len(lines[timing + 1 :]) <= 2
                assert all(len(line) <= 42 for line in lines[timing + 1 :])
    assert render_vtt(Transcript((), "en", "transcribe", "tiny", "cpu", 0)) == b"WEBVTT\r\n\r\n"
